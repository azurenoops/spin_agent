"""Generate local synthetic authorization documents; never call the application."""

from datetime import datetime
from functools import partial
from hashlib import sha256
from io import BytesIO
import json
from math import ceil
from pathlib import Path
from xml.sax.saxutils import escape
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pypdf import PdfReader, PdfWriter
from reportlab.graphics.shapes import Drawing, Line, Rect, String
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.pdfgen.canvas import Canvas
from reportlab.platypus import (
    KeepTogether, LongTable, PageBreak, Paragraph, SimpleDocTemplate, Spacer,
    TableStyle,
)


ROOT = Path(__file__).resolve().parent
NOTICE = "SYNTHETIC TEST DATA - NOT A VALID ATO"
PDFS = {
    "ssp": "01-flankspeed-il5-ssp.pdf",
    "assessment": "02-flankspeed-il5-assessment.pdf",
    "operations": "03-flankspeed-il5-operations.pdf",
    "decision": "04-flankspeed-il5-unsigned-decision.pdf",
}
JSON_NAME = "05-flankspeed-il5-components.json"
WORKBOOK_NAME = "06-flankspeed-il5-registers.xlsx"
NAVY = colors.HexColor("#17324d")
TEAL = colors.HexColor("#176b79")
PALE = colors.HexColor("#eef4f7")
RED = colors.HexColor("#96372d")
WIDTH = 516
STYLES = getSampleStyleSheet()
STYLES.add(ParagraphStyle("ReportTitle", fontName="Helvetica-Bold", fontSize=27,
                         leading=33, textColor=NAVY, spaceAfter=20))
STYLES.add(ParagraphStyle("SectionTitle", fontName="Helvetica-Bold", fontSize=17,
                         leading=22, textColor=NAVY, spaceAfter=12, keepWithNext=True))
STYLES.add(ParagraphStyle("Subhead", fontName="Helvetica-Bold", fontSize=11,
                         leading=15, textColor=TEAL, spaceBefore=10, spaceAfter=5,
                         keepWithNext=True))
STYLES.add(ParagraphStyle("ReportBody", fontName="Helvetica", fontSize=9.5,
                         leading=13, spaceAfter=5, alignment=TA_LEFT))
STYLES.add(ParagraphStyle("Cell", fontName="Helvetica", fontSize=8,
                         leading=11, spaceAfter=0, splitLongWords=True))
STYLES.add(ParagraphStyle("CellHeader", parent=STYLES["Cell"],
                         fontName="Helvetica-Bold", textColor=colors.white))
STYLES.add(ParagraphStyle("Notice", parent=STYLES["ReportBody"], textColor=RED,
                         fontName="Helvetica-Bold", spaceBefore=10, spaceAfter=12))


class SourceHeading(Paragraph):
    def __init__(self, text, source_key, style="SectionTitle"):
        super().__init__(escape(text), STYLES[style])
        self.source_key = source_key
        self.outline_label = text


class SourceDocument(SimpleDocTemplate):
    def __init__(self, filename, label):
        target = str(filename) if isinstance(filename, Path) else filename
        super().__init__(target, pagesize=letter, leftMargin=48,
                         rightMargin=48, topMargin=68, bottomMargin=54,
                         title=f"Flank Speed IL5 - {label} - Synthetic",
                         author="Synthetic fixture generator", subject=NOTICE)
        self.label = label
        self.source_pages = {}

    def afterFlowable(self, flowable):
        if isinstance(flowable, SourceHeading):
            self.source_pages[flowable.source_key] = self.page
            self.canv.bookmarkPage(flowable.source_key)
            self.canv.addOutlineEntry(flowable.outline_label, flowable.source_key, 0)


def text(value, style="ReportBody"):
    return Paragraph(escape(str(value)).replace("\n", "<br/>"), STYLES[style])


def field(label, value):
    return Paragraph(f"<b>{escape(label)}:</b> {escape(str(value)).replace(chr(10), '<br/>')}",
                     STYLES["ReportBody"])


def bullets(values):
    return [text(f"- {value}") for value in values]


def table(headers, rows, widths):
    data = [[text(header, "CellHeader") for header in headers]]
    data.extend([[text(value, "Cell") for value in row] for row in rows])
    result = LongTable(data, colWidths=widths, repeatRows=1, hAlign="LEFT")
    result.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), NAVY),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, PALE]),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LINEBELOW", (0, 0), (-1, 0), 1, TEAL),
        ("LEFTPADDING", (0, 0), (-1, -1), 7),
        ("RIGHTPADDING", (0, 0), (-1, -1), 7),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ]))
    return result


def page_frame(canvas, document):
    canvas.saveState()
    canvas.setFillColor(NAVY)
    canvas.rect(0, 751, 612, 41, fill=1, stroke=0)
    canvas.setFillColor(colors.white)
    canvas.setFont("Helvetica-Bold", 10)
    canvas.drawString(48, 768, "FLANK SPEED IL5")
    canvas.setFont("Helvetica", 8)
    canvas.drawRightString(564, 768, document.label.upper())
    canvas.setFillColor(RED)
    canvas.setFont("Helvetica-Bold", 8)
    canvas.drawString(48, 737, NOTICE)
    canvas.setStrokeColor(TEAL)
    canvas.line(48, 43, 564, 43)
    canvas.setFont("Helvetica", 7)
    canvas.setFillColor(NAVY)
    canvas.drawString(48, 29, "Not the Navy's real package. No real CUI, signature or authorization evidence.")
    canvas.drawRightString(564, 29, f"Page {canvas.getPageNumber()}")
    canvas.restoreState()


def cover(source, title, purpose):
    return [
        Spacer(1, 25), text(title, "ReportTitle"),
        text(source["title"], "SectionTitle"),
        field("Fixture / version", f'{source["fixtureId"]} / {source["version"]}'),
        field("Scenario date", source["asOf"]),
        field("Environment", source["scope"]["environment"]),
        text(NOTICE, "Notice"), text(purpose),
        text("This is an original production-style test scenario. It is not an official "
             "Navy, Microsoft, DISA, eMASS or FedRAMP submission. All operational settings "
             "and findings are invented; public service descriptions do not establish "
             "this tenant's authorization or control effectiveness."),
        table(["Milestone", "Fixture disposition"], [
            ["Source documents", "Generated synthetic reference material"],
            ["Assessment evidence", "Four fictional samples; other evidence explicitly absent"],
            ["Control status", "Illustrative; not assessed"],
            ["Authorization", "Unsigned, not effective; no inheritance verified"],
            ["SPIN workflow", "Receipt, analysis, review, approval and publication remain separate"],
        ], [145, 371]),
        PageBreak(),
    ]


def heading(story, title, key):
    story.append(SourceHeading(title, key))


def boundary_diagram():
    drawing = Drawing(WIDTH, 230)
    lanes = [
        (155, "CUSTOMER / COMMAND - OUTSIDE THIS TENANT BOUNDARY",
         "Managed endpoint | customer network | mission and data owner"),
        (82, "TENANT OPERATOR - NOTIONAL CONFIGURATION BOUNDARY",
         "Identity | access policy | collaboration configuration | operations"),
        (9, "MICROSOFT SERVICE LAYER - CANDIDATE INHERITED DEPENDENCY",
         "Microsoft 365 DoD service infrastructure - evidence not supplied"),
    ]
    for y, title, detail in lanes:
        drawing.add(Rect(0, y, WIDTH, 62, fillColor=PALE, strokeColor=TEAL))
        drawing.add(String(12, y + 40, title, fontName="Helvetica-Bold", fontSize=8.3,
                           fillColor=NAVY))
        drawing.add(String(12, y + 20, detail, fontName="Helvetica", fontSize=8,
                           fillColor=NAVY))
    for y in (154, 81):
        drawing.add(Line(258, y, 258, y - 8, strokeColor=RED))
        drawing.add(Line(258, y - 8, 254, y - 4, strokeColor=RED))
        drawing.add(Line(258, y - 8, 262, y - 4, strokeColor=RED))
    return drawing


def ssp_story(source):
    # Keep the document's ordered sections together so pagination and source
    # anchors can be reviewed against the resulting artifact.
    scope = source["scope"]
    story = cover(source, "System Security Plan",
                  "A tenant-level SSP with explicit shared responsibilities and selected "
                  "control narratives. Forty controls are a representative subset, not a "
                  "complete IL5/FedRAMP/CNSSI baseline or a control-effectiveness claim.")
    heading(story, "1. Document control and authorization boundary", "ssp-overview")
    for label, key in [("Mission", "mission"), ("Tenant operator", "operator"),
                       ("Synthetic tenant identifier", "tenantIdentifier"),
                       ("Categorization assumption", "categorization"), ("Baseline limitation", "baseline")]:
        story.append(field(label, scope[key]))
    story += [text("In-scope logical services", "Subhead"), *bullets(scope["included"]),
              text("Explicit exclusions", "Subhead"), *bullets(scope["excluded"]),
              text("Assumptions requiring validation", "Subhead"), *bullets(scope["assumptions"]),
              PageBreak()]
    heading(story, "2. Boundary and data-flow model", "ssp-boundary")
    story += [boundary_diagram(), Spacer(1, 12),
              text("The diagram shows responsibility layers, not Microsoft's actual internal "
                   "network topology. Arrows are conceptual dependency/data exchanges, not "
                   "approved interconnections or an assurance of encryption."),
              table(["Flow / owner", "Source to destination", "Information / transport", "Boundary / approval"],
                    [[f'{r["id"]}\n{r["owner"]}', f'{r["source"]}\nto {r["destination"]}',
                      f'{r["information"]}\n{r["transport"]}',
                      f'{r["trustBoundary"]}\n{r["approvalStatus"]}'] for r in source["dataFlows"]],
                    [89, 139, 149, 139]), PageBreak()]
    heading(story, "3. Roles and responsibility allocation", "ssp-roles")
    story += [text("All organizations and contacts below are scenario roles, not the real Navy "
                   "chain of command. A CSP administrator is not automatically the mission AO."),
              table(["Role / ID", "Organization / contact", "Assigned responsibilities"],
                    [[f'{r["id"]}\n{r["role"]}', f'{r["organization"]}\n{r["contact"]}',
                      r["responsibilities"]] for r in source["roles"]], [132, 155, 229]),
              PageBreak()]
    heading(story, "4. Component and capability inventory", "ssp-inventory")
    story += [text("Stable identifiers reconcile the repeated PDF, workbook and JSON descriptions. "
                   "Logical operational components are not fabricated servers. No hardware, "
                   "subscription, license or live tenant discovery is asserted."),
              table(["Component", "Type / owner", "Capabilities"],
                    [[f'{r["id"]}\n{r["name"]}', f'{r["type"]}\n{r["owner"]}',
                      ", ".join(r["capabilityIds"])] for r in source["components"]],
                    [222, 169, 125]), PageBreak()]
    capabilities = {r["id"]: r for r in source["capabilities"]}
    for component in source["components"]:
        heading(story, f'{component["id"]} | {component["name"]}', component["id"])
        story += [field("Logical type / owner", f'{component["type"]} / {component["owner"]}'),
                  text(component["description"]),
                  text("Notional configuration assumptions", "Subhead"),
                  *bullets(component["configurationAssumptions"]),
                  text("Tenant/provider responsibilities", "Subhead"),
                  *bullets(component["providerResponsibilities"]),
                  text("Customer retained responsibilities", "Subhead"),
                  *bullets(component["customerResponsibilities"]),
                  field("Evidence references", ", ".join(component["evidenceIds"]))]
        for capability_id in component["capabilityIds"]:
            capability = capabilities[capability_id]
            story += [SourceHeading(f'{capability_id} | {capability["name"]}', capability_id, "Subhead"),
                      text(capability["description"]),
                      field("Proposed controls / contributors", f'{", ".join(capability["controlIds"])} / '
                            f'{", ".join(capability["componentIds"])}'),
                      field("Provider implementation assumption", capability["providerImplementation"]),
                      field("Customer actions", capability["customerActions"])]
        story.append(PageBreak())
    heading(story, "5. Control implementation and assessment crosswalk", "ssp-controls")
    story += [text("The following forty parent controls are selected examples. Enhancements, "
                   "organization-defined parameters, privacy/NSS overlays and service-level "
                   "inheritance require a real tailored baseline. No omitted control is "
                   "implicitly satisfied or not applicable."),
              table(["Control", "Responsibility", "Open finding references"],
                    [[f'{r["id"]}\n{r["title"]}', r["ownership"],
                      ", ".join(r["gapFindings"]) or "No finding modeled; not evidence of effectiveness"]
                     for r in source["controls"]], [220, 141, 155]), PageBreak()]
    for control in source["controls"]:
        heading(story, f'{control["id"]} | {control["title"]}', control["id"])
        story += [field("Assessment status", control["status"]), field("Allocation", control["ownership"]),
                  text(control["narrative"]),
                  text("Notional organization-defined parameters", "Subhead"), text(control["parameters"]),
                  text("Examine / interview / test", "Subhead"), text(control["assessmentProcedure"]),
                  text("Acceptance criterion", "Subhead"), text(control["acceptanceCriterion"]),
                  field("Evidence", ", ".join(control["evidenceIds"])),
                  field("Known modeled gaps", ", ".join(control["gapFindings"]) or "No modeled finding; not assessed"),
                  text("Acceptance criteria are future assessment conditions, not recorded test "
                       "results. Refer to the evidence register before claiming inheritance."),
                  PageBreak()]
    heading(story, "6. Limitations and public reference basis", "ssp-references")
    story += bullets(source["limitations"])
    for reference in source["references"]:
        story += [text(f'{reference["id"]} | {reference["title"]}', "Subhead"),
                  text(reference["url"]), field("Read on", reference["checkedOn"]),
                  text(reference["applicability"]),
                  text("Public background only; not tenant authorization evidence.")]
    return story


def assessment_story(source):
    story = cover(source, "Assessment Plan, Report and Remediation",
                  "A production-style SAP/SAR/risk/POA&M companion. No actual assessment "
                  "was performed. Six invented open findings are evidence-gap scenarios, "
                  "not allegations about the Navy or Microsoft's operating service.")
    heading(story, "1. Assessment plan and rules of engagement", "sap")
    story += [field("Scope", source["scope"]["mission"]),
              text("Methods follow the examine/interview/test structure described by NIST "
                   "SP 800-53A. The forty SSP controls are a representative exercise scope. "
                   "Sampling and thresholds are notional local targets, not DoD mandates."),
              table(["Stage", "Required activity", "Exit condition"], [
                  ["Preparation", "Agree the exact tenant/boundary version, assessor independence, "
                   "licensed service inventory, evidence handling and named approvals.",
                   "Approved plan and rules of engagement; none issued in this fixture."],
                  ["Examine", "Review dated exports, role grants, change records, logging coverage, "
                   "retention decisions and current provider inheritance evidence.",
                   "Each evidence item has identity, date, scope, custodian and integrity record."],
                  ["Interview", "Interview operator, command data owner, SOC, privacy and continuity roles. "
                   "Reconcile inconsistent descriptions with recorded configurations.",
                   "Reviewed interview notes and explicit unresolved questions."],
                  ["Test", "Use isolated test identities and invented records. Check allowed/denied "
                   "access, revocation, approved sharing, telemetry delivery and scoped restore.",
                   "Expected/observed results and repeatable evidence; no production testing here."],
                  ["Report", "Record limitations, unmet criteria, source severity and remediation.",
                   "Independent assessor review; findings remain open until evidence is accepted."],
              ], [76, 274, 166]),
              text("Prohibited activity", "Subhead"),
              text("No live scanning, account creation, privilege escalation, tenant changes or "
                   "external file transfer is authorized by these documents. No real PII or CUI "
                   "may be introduced into this local fixture."),
              PageBreak()]
    heading(story, "2. Assessment results and risk summary", "sar")
    story += [text("Illustrative outcome: 2 High, 3 Moderate and 1 Low open findings. "
                   "These source-stated labels exercise review and prioritization; no real "
                   "risk score, risk acceptance, control pass or ATO recommendation is issued."),
              table(["Finding", "Severity / status", "Controls", "Evidence references"],
                    [[f'{r["id"]}\n{r["title"]}', f'{r["severity"]}\n{r["status"]}',
                      ", ".join(r["controls"]), ", ".join(r["evidenceIds"])]
                     for r in source["findings"]], [235, 80, 83, 118]),
              text("All other controls remain illustrative and unassessed, not implicitly "
                   "passed. Missing provider evidence is an inheritance-verification gap, "
                   "not proof that the provider's control failed."),
              PageBreak()]
    for finding in source["findings"]:
        heading(story, f'{finding["id"]} | {finding["title"]}', finding["id"])
        story += [field("Source severity / workflow state", f'{finding["severity"]} / {finding["status"]}'),
                  field("Controls / components", f'{", ".join(finding["controls"])} / '
                        f'{", ".join(finding["componentIds"])}'),
                  text("Illustrative observation", "Subhead"), text(finding["observation"]),
                  text("Potential impact in this scenario", "Subhead"), text(finding["impact"]),
                  text("Recommended response", "Subhead"), text(finding["recommendation"]),
                  field("Evidence references", ", ".join(finding["evidenceIds"])),
                  text("Assessment date: 2026-09-24 (fictional scenario). Assessor: Synthetic "
                       "independent assessor role. No actual service testing occurred."),
                  PageBreak()]
    for poam in source["poams"]:
        heading(story, f'{poam["id"]} | Remediation plan', poam["id"])
        story += [field("Linked finding", poam["findingId"]), field("Owner", poam["owner"]),
                  field("Workflow status", poam["status"]), text(poam["correctiveAction"]),
                  field("Resources / dependencies", poam["resources"]),
                  field("Scheduled completion (notional)", poam["scheduledCompletion"]),
                  table(["Milestone", "Due date"], [[r["description"], r["dueDate"]]
                        for r in poam["milestones"]], [405, 111]),
                  text("Required closure evidence", "Subhead"), *bullets(poam["closureEvidence"]),
                  text("Submitted closure evidence: none. Evidence upload, source text and due "
                       "dates never close a finding. A separate authorized reviewer must "
                       "evaluate the exact finding revision and supporting evidence."),
                  PageBreak()]
    heading(story, "3. Assessment limitations and AO handoff", "sar-handoff")
    story += [text("This package supplies scenario records and sample artifacts only. The "
                   "real assessor would reconcile all controls/enhancements, verify current "
                   "provider documents, observe tests, identify unsupported services and "
                   "report residual risk. The AO decides authorization; the tool does not."),
              *bullets(source["limitations"])]
    return story


def operations_story(source):
    story = cover(source, "Operational Annexes and Evidence Register",
                  "Tenant operations, continuous monitoring and evidence handling with "
                  "specific notional handoffs. These are proposed exercise procedures, "
                  "not proof that Navy or Microsoft operations follow these settings.")
    for procedure in source["procedures"]:
        heading(story, f'{procedure["id"]} | {procedure["title"]}', procedure["id"])
        story += [field("Accountable role", procedure["owner"])]
        for index, step in enumerate(procedure["steps"], 1):
            story.append(text(f"{index}. {step}"))
        story += [field("Exit / handoff criteria", procedure["exitCriteria"]),
                  field("Limitations", procedure["limitations"]), PageBreak()]
    heading(story, "Continuous monitoring strategy", "conmon")
    story += [text("All frequencies, thresholds and targets are fictional policy selections. "
                   "Actual contractual obligations and approved DoD reporting channels take "
                   "precedence. Do not conflate service health with control effectiveness.")]
    for row in source["monitoring"]:
        story.append(KeepTogether([
            text(f'{row["id"]} | {row["activity"]}', "Subhead"),
            field("Frequency / owner", f'{row["frequency"]} / {row["owner"]}'),
            field("Measure / threshold", f'{row["measure"]} / {row["threshold"]}'),
            field("Action / evidence", f'{row["action"]} / {", ".join(row["evidenceIds"])}'),
        ]))
    story.append(PageBreak())
    heading(story, "Evidence inventory and custody", "evidence-index")
    story += [text("Four original fictional examples are printed in this annex. Twelve required "
                   "evidence items are absent. Describing a missing item is not supplying it. "
                   "Public links are background references, not controlled provider artifacts."),
              table(["Evidence / owner", "State / source type", "Purpose / review cycle"],
                    [[f'{r["id"]}\n{r["title"]}\n{r["owner"]}', f'{r["state"]}\n{r["sourceType"]}',
                      f'{r["description"]}\nReview: {r["reviewCycle"]}'] for r in source["evidence"]],
                    [173, 132, 211]), PageBreak()]
    for evidence in source["evidence"]:
        if evidence["state"] != "Synthetic sample included":
            continue
        heading(story, f'{evidence["id"]} | {evidence["title"]}', evidence["id"])
        story += [text("FICTIONAL SAMPLE - generated locally, not collected from a tenant.", "Notice"),
                  field("Owner / type", f'{evidence["owner"]} / {evidence["sourceType"]}'),
                  field("Controls / components", f'{", ".join(evidence["relatedControls"])} / '
                        f'{", ".join(evidence["relatedComponents"])}'),
                  text(evidence["sampleText"]),
                  text("Custody: local synthetic generation only. A real evidence record "
                       "requires collection time, collector, approved storage, hash and "
                       "scope/revision binding. No closure or acceptance is implied."),
                  PageBreak()]
    heading(story, "Production-readiness evidence gaps", "operations-gaps")
    story += [text("Before a real submission, obtain controlled provider authorization and "
                   "customer-responsibility artifacts through approved channels; reconcile "
                   "actual tenant exports and feature availability; complete baseline tailoring, "
                   "privacy/interconnection reviews, independent assessment and AO decisions."),
              *bullets(source["limitations"])]
    return story


def decision_story(source):
    decision = source["decision"]
    story = cover(source, "Unsigned Authorization Decision Draft",
                  "A decision-reference exercise, not an authorization instrument. No official "
                  "signature, seal, eMASS identifier, valid effective date or verified provider "
                  "authorization is supplied.")
    heading(story, "Proposed decision context - not issued", decision["reference"])
    for label, key in [("Reference", "reference"), ("Authority placeholder", "authority"),
                       ("Decision type", "decisionType"), ("Actual status", "status"),
                       ("Proposed start only", "proposedStartDate"),
                       ("Proposed expiration only", "proposedExpirationDate"), ("Scope", "scope")]:
        story.append(field(label, decision[key]))
    story += [field("Actual issuance / effective date", "Not established"),
              field("Signature / verification", "Absent; no official has reviewed or signed this draft"),
              text("Proposed conditions", "Subhead"), *bullets(decision["conditions"]),
              text("Qualifications", "Subhead"), *bullets(decision["qualifications"]),
              text("No mission system ATO, covered-workload finding, inherited control acceptance "
                   "or reusable capability publication is granted by importing this document. "
                   "Proposed dates must not be stored as the dates of an issued authorization.")]
    heading(story, "Decision review and release checklist", "decision-checklist")
    story += [table(["Required check", "This fixture"], [
        ["Actual authorizing official and authority verified", "Not supplied"],
        ["Current provider scope and inheritance evidence reviewed", "Not supplied"],
        ["System categorization and complete tailored baseline approved", "Representative subset only"],
        ["Independent assessment and closure/accepted residual risk", "Six fictional open findings; no real assessment"],
        ["Exact boundary, dependencies and impact preview reviewed", "Requires application/operator review"],
        ["Signed and dated decision with limitations and expiration", "Unsigned draft; no effective authorization"],
        ["Capability publication separated from mission authorization", "Required; never inferred from upload or confidence"],
    ], [317, 199]), text(NOTICE, "Notice")]
    return story


def render_pdf(name, label, story):
    document = SourceDocument(ROOT / name, label)
    document.build(story, onFirstPage=page_frame, onLaterPages=page_frame,
                   canvasmaker=partial(Canvas, invariant=1))
    return document.source_pages


def projection(source, pages):
    qualification = [NOTICE, "Original fictional scenario; no actual Navy/Microsoft service assessed.",
                     "Human review is required. No authorization, closure or publication is granted."]
    source_note = lambda key: f" Source: {PDFS['ssp']}, page {pages['ssp'][key]}, {key}."
    result = {
        "components": [{"id": r["id"], "type": r["type"], "name": r["name"],
                        "description": f'{NOTICE}. {r["description"]}{source_note(r["id"])}'}
                       for r in source["components"]],
        "capabilities": [{
            "id": r["id"], "name": r["name"], "type": "Service",
            "description": f'{NOTICE}. {r["description"]} Provider: {r["providerImplementation"]} '
                           f'Customer: {r["customerActions"]}{source_note(r["id"])}',
            "componentIds": r["componentIds"], "controlIds": r["controlIds"], "responsibility": "Shared",
        } for r in source["capabilities"]],
    }
    decision = source["decision"]
    result["authorizationDecisionClaims"] = [{
        "kind": "AuthorizationDecisionClaim", "id": decision["reference"],
        "name": "Flank Speed IL5 - unsigned synthetic decision draft",
        "claim": {
            "authorizationDecision": {
                "subjectKind": "Provider", "subject": source["title"], "reference": decision["reference"],
                "authority": decision["authority"], "decisionType": decision["decisionType"],
                "statusAsStated": decision["status"], "scope": decision["scope"],
                "conditions": decision["conditions"], "exclusions": source["scope"]["excluded"],
            },
            "relationships": [{"kind": "DecisionBoundary", "targetSourceId": f"FS-BOUNDARY-{kind}"}
                              for kind in ("INCLUDED", "EXCLUDED")],
            "qualifications": qualification + decision["qualifications"],
        },
    }]
    result["boundaryClaims"] = [{
        "kind": "BoundaryClaim", "id": f"FS-BOUNDARY-{kind.upper()}",
        "name": f"{kind} Flank Speed IL5 notional scope",
        "claim": {
            "boundary": {"subject": source["title"], "relationship": kind,
                         "scope": "; ".join(source["scope"][kind.lower()]),
                         "environment": source["scope"]["environment"],
                         "responsibilities": [
                             "Tenant operator controls the fictional configuration boundary.",
                             "Microsoft common controls are candidate inheritance, not verified evidence.",
                             "Customers retain mission authorization, endpoints, networks and data ownership.",
                         ], "decisionReference": decision["reference"]},
            "relationships": [{"kind": "BoundaryComponent", "targetSourceId": r["id"]}
                              for r in source["components"]] if kind == "Included" else [],
            "qualifications": qualification,
        },
    } for kind in ("Included", "Excluded")]
    result["assessmentFindings"] = [{
        "kind": "AssessmentFinding", "id": r["id"], "name": f'{r["id"]} | {r["title"]}',
        "claim": {
            "assessmentFinding": {
                "sourceFindingId": r["id"], "observation": f'{NOTICE}. {r["observation"]}',
                "severityAsStated": r["severity"], "statusAsStated": "Open",
                "assessor": "Synthetic independent assessor role", "assessmentDate": source["asOf"],
                "controlIds": r["controls"],
            },
            "relationships": [{"kind": "FindingComponent", "targetSourceId": component}
                              for component in r["componentIds"]],
            "qualifications": qualification,
        },
    } for r in source["findings"]]
    result["poamItems"] = [{
        "kind": "PoamItem", "id": r["id"], "name": r["id"],
        "claim": {
            "poamItem": {"sourcePoamId": r["id"], "correctiveAction": r["correctiveAction"],
                         "ownerAsStated": r["owner"], "statusAsStated": "Open",
                         "milestones": r["milestones"], "requiredClosureEvidence": r["closureEvidence"],
                         "submittedEvidenceReferences": []},
            "relationships": [{"kind": "PoamFinding", "targetSourceId": r["findingId"]}],
            "qualifications": qualification,
        },
    } for r in source["poams"]]
    return result


def flat(value):
    if isinstance(value, list):
        return "\n".join(flat(item) for item in value)
    if isinstance(value, dict):
        return "; ".join(f"{key}: {flat(item)}" for key, item in value.items())
    return str(value)


def create_workbook(source):
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.properties.title = source["title"]
    workbook.properties.creator = "Synthetic fixture generator"
    workbook.properties.description = NOTICE
    workbook.properties.created = datetime(2026, 9, 24)
    workbook.properties.modified = datetime(2026, 9, 24)
    sheets = {
        "Metadata": [{"Field": "Fixture", "Value": source["fixtureId"]},
                     {"Field": "Status", "Value": NOTICE},
                     {"Field": "Scope", "Value": source["scope"]["environment"]},
                     {"Field": "Baseline", "Value": source["scope"]["baseline"]},
                     {"Field": "Decision", "Value": source["decision"]["status"]}],
        "Components": source["components"], "Capabilities": source["capabilities"],
        "Controls": source["controls"], "Findings": source["findings"], "POAM": source["poams"],
        "Evidence": source["evidence"], "Roles": source["roles"], "DataFlows": source["dataFlows"],
        "Monitoring": source["monitoring"], "Procedures": source["procedures"],
        "References": source["references"],
    }
    for name, rows in sheets.items():
        sheet = workbook.create_sheet(name)
        columns = list(rows[0]) + ["FixtureNotice"]
        sheet.append(columns)
        for record in rows:
            sheet.append([flat(record.get(key, "")) for key in columns[:-1]] + [NOTICE])
        for row in sheet:
            for cell in row:
                cell.data_type = "s"
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                cell.font = Font(name="Calibri", size=11)
                if cell.row == 1:
                    cell.fill = PatternFill("solid", fgColor="17324D")
                    cell.font = Font(name="Calibri", size=11, color="FFFFFF", bold=True)
                elif cell.row % 2:
                    cell.fill = PatternFill("solid", fgColor="EEF4F7")
        for index, key in enumerate(columns, 1):
            sheet.column_dimensions[get_column_letter(index)].width = (
                20 if key.lower() in ("id", "status", "severity") else 62 if key in
                ("narrative", "description", "observation", "assessmentProcedure", "sampleText", "steps") else 38)
        sheet.row_dimensions[1].height = 30
        for row in sheet.iter_rows(min_row=2):
            wrapped_lines = max(
                sum(max(1, ceil(len(line) / (sheet.column_dimensions[cell.column_letter].width - 3)))
                    for line in str(cell.value or "").split("\n"))
                for cell in row
            )
            sheet.row_dimensions[row[0].row].height = min(400, max(30, wrapped_lines * 15 + 6))
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.print_title_rows = "1:1"
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.oddFooter.center.text = NOTICE
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    with ZipFile(BytesIO(stream.getvalue())) as original:
        content = {name: original.read(name) for name in original.namelist()}
    write_archive(ROOT / WORKBOOK_NAME, content)


def write_json(name, value):
    (ROOT / name).write_text(json.dumps(value, indent=2, ensure_ascii=True) + "\n", encoding="utf-8")


def write_archive(path, content):
    with ZipFile(path, "w", compression=ZIP_DEFLATED) as archive:
        for name, data in content.items():
            info = ZipInfo(name, date_time=(2026, 9, 24, 0, 0, 0))
            info.compress_type = ZIP_DEFLATED
            info.external_attr = 0o644 << 16
            archive.writestr(info, data)


def locked_appendix():
    stream = BytesIO()
    document = SourceDocument(stream, "Encrypted exception")
    document.build([text(NOTICE, "Notice"), text("Locked synthetic appendix"),
                    text("This deliberately encrypted test attachment must be reported unreadable "
                         "or handled through an explicitly supported decryption workflow.")],
                   onFirstPage=page_frame, canvasmaker=partial(Canvas, invariant=1))
    reader = PdfReader(BytesIO(stream.getvalue()))
    writer = PdfWriter()
    for page in reader.pages:
        writer.add_page(page)
    writer.encrypt("fixture-only-password", algorithm="AES-256")
    output = BytesIO()
    writer.write(output)
    return output.getvalue()


def expected_results(source, pages):
    expected = {
        "fixture": source["fixtureId"], "synthetic": True,
        "expectationsType": "target-workflow acceptance criteria, not a recorded application result",
        "uniqueComponentCount": len(source["components"]),
        "uniqueCapabilityCount": len(source["capabilities"]),
        "notes": [
            "PDF, workbook and JSON repeat the same identities; do not multiply records.",
            "Forty control narratives are a representative subset, not a complete IL5 baseline.",
            "Four fictional evidence samples are included; twelve requested artifacts are absent.",
            "No authorization is effective. Proposed draft dates are not issued-decision dates.",
            "Findings and POAMs are not capabilities, and imported plans do not close findings.",
            "All generated inventory remains unpublished pending human review and exact approval.",
            "Semantic extraction and application review may be partial; structural checks are not acceptance.",
        ],
        "attentionEntries": [
            {"path": "07-encrypted-appendix.pdf", "expected": "Unreadable",
             "reason": "Actually encrypted with a public fixture password; do not silently skip."},
            {"path": "08-malformed-source.json", "expected": "Failed",
             "reason": "Intentionally invalid JSON syntax; not a valid source with zero candidates."},
            {"path": "09-unsupported-appendix.synthetic", "expected": "Unsupported",
             "reason": "Deliberately unsupported extension; retain and account for the entry."},
        ],
    }
    for key in ("components", "capabilities", "controls", "findings", "poams"):
        document = "assessment" if key in ("findings", "poams") else "ssp"
        expected[key] = [{"fixtureId": row["id"], "name": row.get("name", row.get("title", row["id"])),
                          "source": {"artifact": PDFS[document], "page": pages[document][row["id"]],
                                     "section": row["id"]}}
                         for row in source[key]]
    return expected


def main():
    source = json.loads((ROOT / "source-content.json").read_text(encoding="utf-8"))
    pages = {}
    for key, label, story_builder in [
        ("ssp", "System Security Plan", ssp_story),
        ("assessment", "Assessment and Remediation", assessment_story),
        ("operations", "Operations and Evidence", operations_story),
        ("decision", "Unsigned Decision Draft", decision_story),
    ]:
        pages[key] = render_pdf(PDFS[key], label, story_builder(source))
    write_json(JSON_NAME, projection(source, pages))
    create_workbook(source)
    names = [*PDFS.values(), JSON_NAME, WORKBOOK_NAME]
    clean = {name: (ROOT / name).read_bytes() for name in names}
    attention = {
        **clean, "07-encrypted-appendix.pdf": locked_appendix(),
        "08-malformed-source.json": b'{"synthetic":true,"components":[INVALID JSON',
        "09-unsupported-appendix.synthetic": (NOTICE + "\nUnsupported format fixture; not a signature.").encode(),
    }
    write_archive(ROOT / "flankspeed-il5-ato-clean.zip", clean)
    write_archive(ROOT / "flankspeed-il5-ato-needs-attention.zip", attention)
    write_json("expected-results.json", expected_results(source, pages))
    manifest = {"schemaVersion": 2, "fixture": source["fixtureId"], "synthetic": True,
                "sourceSha256": sha256((ROOT / "source-content.json").read_bytes()).hexdigest(),
                "artifacts": {}, "archives": {}}
    for name, content in clean.items():
        entry = {"sha256": sha256(content).hexdigest(), "bytes": len(content)}
        if name.endswith(".pdf"):
            entry["pages"] = len(PdfReader(BytesIO(content)).pages)
        manifest["artifacts"][name] = entry
    for name, content in [("flankspeed-il5-ato-clean.zip", clean),
                          ("flankspeed-il5-ato-needs-attention.zip", attention)]:
        data = (ROOT / name).read_bytes()
        manifest["archives"][name] = {"sha256": sha256(data).hexdigest(), "bytes": len(data),
                                      "entries": list(content)}
    write_json("package-manifest.json", manifest)
    print(json.dumps({"artifacts": manifest["artifacts"],
                      "archiveEntryCounts": {key: len(value["entries"])
                                            for key, value in manifest["archives"].items()}}, indent=2))


if __name__ == "__main__":
    main()
